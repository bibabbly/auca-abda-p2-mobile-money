"""Build the Project 2 report as a Word document in milestone order.

Screenshots come from screenshots/, charts from results/charts/, diagrams from
reports/figures/ (run reports/diagrams.py first). Timing and load-test tables are
read from results/ so they match the raw data. Missing screenshots are marked in red.

Usage:  .venv\\Scripts\\python.exe reports\\build_report.py
Output: reports/ABDA_Project2_Report.docx  (edit in Word, then save as PDF)
"""
import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS, CHARTS, FIGS, RES = ROOT / "screenshots", ROOT / "results" / "charts", ROOT / "reports" / "figures", ROOT / "results"
OUT = ROOT / "reports" / "ABDA_Project2_Report.docx"
WIDTH = Cm(17)
RED = RGBColor(0xC0, 0x00, 0x00)

doc = Document()
missing = []
fig_no = [0]


# ------------------------------------------------------------------ helpers
def setup():
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    footer = sec.footer.paragraphs[0]
    footer.text = "AUCA Advanced Big Data Analytics, Project 2: Mobile Money Investigations | Page "
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)


def h(text, level=1):
    doc.add_heading(text, level=level)


def p(text, bold=False, italic=False, color=None, size=None, align=None):
    para = doc.add_paragraph()
    for i, chunk in enumerate(text.split("**")):
        run = para.add_run(chunk)
        run.bold = bold or (i % 2 == 1)
        run.italic = italic
        if color:
            run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)
    if align:
        para.alignment = align
    return para


def bullets(items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        for i, chunk in enumerate(it.split("**")):
            para.add_run(chunk).bold = i % 2 == 1


def todo(text):
    p(f"[TO DO] {text}", bold=True, color=RED)


def table(header, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = ""
        r = cell.paragraphs[0].add_run(str(text)); r.bold = True; r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(text)); r.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Cm(w)
    doc.add_paragraph()
    return t


def img(path, caption, width=WIDTH):
    fig_no[0] += 1
    if not Path(path).exists():
        missing.append(Path(path).name)
        todo(f"Insert screenshot {Path(path).name}: {caption}")
        return
    doc.add_picture(str(path), width=width)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p(f"Figure {fig_no[0]}. {caption}", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)


def shot(name, caption, width=WIDTH):
    img(SHOTS / f"{name}.png", caption, width)


def read_csv(path):
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def ms(v):
    return "-" if v in ("", None) else f"{float(v):,.1f}"


# ------------------------------------------------------------------ content
def cover():
    for _ in range(4):
        doc.add_paragraph()
    p("Advanced Big Data Analytics", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    p("Project 2: Mobile Money Investigations with PostgreSQL and Neo4j", bold=True, size=20,
      align=WD_ALIGN_PARAGRAPH.CENTER)
    p("Adventist University of Central Africa (AUCA), Kigali", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    table(["Member", "Student ID"], [["BIZIMUNGU Theoneste", "101398"], ["NIYITANGA Eric", "101393"],
                                     ["NDASHIMYE Janvier", "101363"]], widths=[8, 4], size=11)
    p("Submitted 26 October 2026", align=WD_ALIGN_PARAGRAPH.CENTER)
    p("Code and raw results: https://github.com/bibabbly/auca-abda-p2-mobile-money",
      align=WD_ALIGN_PARAGRAPH.CENTER)
    todo("Check names, spelling and student IDs; remove this line.")
    doc.add_page_break()

    h("Summary", 1)
    bullets([
        "September 2026 data (20,993,500 records, 20,003,300 transfers) captured from the unchanged generator, "
        "loaded into PostgreSQL 18.4 and Neo4j 2026.10.1 Community; all counts and audit totals reconcile.",
        "Q1–Q7 implemented as SQL/Cypher pairs; **all 61 cases return identical answers** in both databases, "
        "confirmed again by two checks recomputed directly from the CSV files.",
        "**PostgreSQL** is faster for indexed lookups (Q1/Q3/Q5/Q6: 0.4–3 ms vs 5.5–8.4 ms). "
        "**Neo4j** is faster for deep traversals (Q4 depth 10: 4–6 ms vs 23 ms; Q7: 8–14 ms).",
        "Plain recursive SQL for Q7 bound 6 timed out (>60 s); a **bidirectional SQL rewrite** gives the same "
        "answers in ≈80 ms, so most of that gap was the algorithm, not the model.",
        "Under concurrent load both engines saturate the 8 CPU threads at 8 clients; Neo4j delivered "
        "12–32× higher throughput for Q4/Q7 on this machine.",
        "Investigate **A (998), B (147228) and C (206978)** before D, E, F: same device and network profile, "
        "but only A–C show rapid onward chains and/or circular payments.",
    ])
    doc.add_page_break()


def section1():
    h("1 Capture the data", 1)
    h("1.1 Workspace", 2)
    table(["Item", "Value"], [
        ["CPU / RAM / OS", "Intel Core i7-1165G7 (4 cores, 8 threads) / 34.1 GB / Windows 11 Pro 10.0.22631"],
        ["PostgreSQL", "18.4, database mobile_money, TimeZone Africa/Kigali; shared_buffers 8 GB, "
                       "effective_cache_size 18 GB, work_mem 64 MB, maintenance_work_mem 2 GB, max_wal_size 8 GB"],
        ["Neo4j", "2026.10.1 Community, default neo4j database; heap 8 GB, page cache 10 GB, "
                  "transaction memory 6 GB; Java 21.0.11"],
        ["Python / drivers", "Python 3.13.6; psycopg 3.3.6 (libpq 18.0.4); neo4j 6.4.1; numpy 2.5.4; pandas 3.1.0; "
                             "matplotlib 3.11.2"],
        ["Generator", "mobile_money_generator.py unchanged, SHA-256 2ea0e59d…9455a"],
    ], widths=[4, 13])
    shot("1.1_postgres", "PostgreSQL database created with time zone Africa/Kigali.")
    shot("1.1_environment_0", "env_check.py: machine, drivers, PostgreSQL settings.")
    shot("1.1_environment_1", "env_check.py: Neo4j version and memory settings; both connections succeed.")

    h("1.2 capture.py", 2)
    p("capture.py imports the generator as g, previews ten records per stream with itertools.islice, and streams "
      "each iterator straight into csv.DictWriter, counting consumed and written rows. The full stream is never held "
      "in memory. Transfer audit totals are accumulated while writing.")
    audit = json.loads((RES / "capture_audit.json").read_text(encoding="utf-8"))
    rows = [[k, f"{v['written']:,}", f"{v['expected']:,}", "OK" if v["matches"] else "MISMATCH",
             f"{v['capture_seconds']:.1f}", f"{v['csv_gb']:.3f}"] for k, v in audit["streams"].items()]
    rows.append(["TOTAL", f"{audit['total_rows']:,}", "20,993,500", "OK" if audit["all_match"] else "MISMATCH",
                 f"{audit['total_capture_seconds']:.1f}", f"{audit['total_csv_gb']:.3f}"])
    table(["Stream", "Written", "Expected", "Status", "Seconds", "CSV GB"], rows)
    a = audit["streams"]["transfers"]["audit"]
    p(f"Transfer audit: count {a['count']:,}; earliest {a['earliest_occurred_at']} ({a['earliest_cat'][:19]} CAT); "
      f"latest {a['latest_occurred_at']} ({a['latest_cat'][:19]} CAT); sum amount_rwf {a['sum_amount_rwf']:,} RWF.")
    shot("1.2_preview_0", "Preview: dataset_info() and FIELDS.")
    shot("1.2_preview_1", "Preview: first ten records of each stream.")
    shot("1.2_capture_totals", "Completed five-stream capture with counts, timings, CSV sizes and audit totals.")
    doc.add_page_break()


def section2():
    h("2 Build PostgreSQL", 1)
    h("2.1 Five tables", 2)
    img(FIGS / "pg_schema.png", "PostgreSQL schema: primary keys, foreign keys and checks.")
    table(["Field", "Meaning and unit"], [
        ["customer_id / region", "Customer identifier; home region (Kigali, Northern, Southern, Eastern, Western)"],
        ["account_id / role", "Wallet identifier; person (1–240,000), merchant (240,001–246,000), agent (246,001–248,000)"],
        ["device_id / access_type", "Device identifier; handset or public_terminal"],
        ["first_day / last_day", "Inclusive observation days, 0 = 1 Sep 2026 CAT, 29 = 30 Sep (SMALLINT)"],
        ["observations", "Count of sightings of the account on the device within the interval"],
        ["transfer_id", "Unique payment event (BIGINT)"],
        ["sender_id / receiver_id", "Paying and receiving wallets (FKs to accounts)"],
        ["occurred_at", "Unix epoch seconds (BIGINT); displayed in CAT = UTC+02:00"],
        ["amount_rwf", "Whole Rwandan francs, > 0 (INTEGER)"],
    ], widths=[4.5, 12.5])
    shot("2.1_definitions_0", "Table definitions (\\d).")
    shot("2.1_definitions_1", "Table definitions, indexes and foreign keys.")
    shot("2.1_definitions_2", "All constraints validated.")

    h("2.2 Bulk load and checks", 2)
    p("Tables were created with NOT NULL columns only, loaded with \\copy (customers and devices, then accounts, "
      "account_devices, transfers), and only then given keys, foreign keys, checks, indexes and ANALYZE, so load "
      "and setup times are separate.")
    table(["Step", "Seconds"], [["\\copy load (5 files)", "34.6"], ["Keys, FKs, checks, 5 indexes, ANALYZE", "96.8"],
                                ["Total", "131.4"]], widths=[10, 4])
    pg = json.loads((RES / "pg_checks.json").read_text(encoding="utf-8"))
    s = pg["sizes"]
    p(f"Counts, transfer audit totals and roles (240,000 / 6,000 / 2,000) all match capture.py; 13 constraints "
      f"validated; 17/17 columns NOT NULL. Size: data {s['_tables_total']['data_bytes'] / 1e9:.3f} GB, indexes "
      f"{s['_tables_total']['index_bytes'] / 1e9:.3f} GB, database {s['_database_bytes'] / 1e9:.3f} GB "
      f"(1 GB = 10^9 bytes). The two covering transfer indexes are ≈ 0.99 GB each.")
    shot("2.2_load", "\\copy load with per-file timings: 34.6 s.")
    shot("2.2_setup", "Constraints, indexes and ANALYZE: 96.8 s.")
    shot("2.2_checks_0", "pg_checks.py: counts, audit totals, roles, constraints, indexes.")
    shot("2.2_checks_1", "pg_checks.py: data/index sizes and PASSED.")
    doc.add_page_break()


def section3():
    h("3 Build Neo4j", 1)
    h("3.1 Graph model", 2)
    img(FIGS / "neo_schema.png", "Graph schema: Customer, Account, Device nodes; OWNS, USED_DEVICE, TRANSFER.")
    table(["Source field", "Graph element"], [
        ["customers.customer_id, region", ":Customer {customer_id, region}"],
        ["accounts.account_id, role", ":Account {account_id, role}"],
        ["accounts.customer_id", "(:Customer)-[:OWNS]->(:Account) (ignored on the node)"],
        ["devices.device_id, access_type", ":Device {device_id, access_type}"],
        ["account_devices.*", "(:Account)-[:USED_DEVICE {first_day, last_day, observations}]->(:Device)"],
        ["transfers.*", "(:Account)-[:TRANSFER {transfer_id, occurred_at, amount_rwf}]->(:Account), one per row"],
    ], widths=[6, 11])

    h("3.2 Import", 2)
    st = json.loads((RES / "neo4j_setup.json").read_text(encoding="utf-8"))
    p("prepare_import.py writes temporary copies with typed headers (ID groups Customer, Account, Device; :int/:long "
      "numbers) and derives owns.csv in (customer_id, account_id) order; the source CSVs are unchanged. "
      "neo4j-admin database import full ran with --id-type=integer, three node and three relationship files.")
    table(["Measure", "Value"], [
        ["Offline import", f"{st['import_seconds']} s"], ["Constraints + range index online", f"{st['setup_seconds']} s"],
        ["Store size", f"{st['store_gb']:.3f} GB"], ["Transaction logs", f"{st['transaction_log_gb']:.3f} GB"],
        ["Nodes / relationships", "704,600 / 20,536,900 (all checks passed)"],
    ], widths=[8, 8])
    shot("3.2_import", "neo4j-admin full import completed.")
    shot("3.2_setup", "Unique constraints and TRANSFER.occurred_at range index online; store and log sizes.")
    shot("3.2_checks", "Node/relationship counts, roles, one owner per account, unique transfer_id, audit totals.")
    shot("3.2_clean", "Temporary import copies removed.")

    h("3.3 One account in both models", 2)
    shot("3.3_inspect", "Account 998 traced to its owner, a device and an outgoing transfer: identical in SQL and Cypher.")
    p("A foreign key stores the other row's key and must be joined at query time; in the graph the same fact is a "
      "stored relationship, so accounts.customer_id becomes OWNS and account_devices becomes USED_DEVICE with its "
      "interval as properties. If repeated payments between two accounts were collapsed into one edge, each payment's "
      "transfer_id, occurred_at and amount_rwf would be lost (or reduced to a count/sum), so Q1 hourly totals, the "
      "time-ordered chains of Q5 and the cycles of Q6 could no longer be evaluated.")
    doc.add_page_break()


def section4():
    h("4 Routine business questions", 1)
    p("Review hour: 09:45:00 inclusive to 10:45:00 exclusive CAT on each set's date, converted to epoch seconds in "
      "Python. Sets: A 998 (6 Sep), B 147228 (13 Sep), C 206978 (8 Sep), D 237998 (6 Sep), E 144228 (13 Sep), "
      "F 203978 (8 Sep).")
    h("4.1 Q1 Account activity", 2)
    table(["Set", "A", "B", "C", "D", "E", "F"],
          [["Count", 1, 1, 1, 1, 1, 1],
           ["Total RWF", "200,000", "1,050,000", "300,000", "45,000", "40,000", "30,000"]])
    shot("4.1_q1", "Q1 SQL and Cypher results for all six sets (identical).")
    p("Outgoing value is not a balance: a balance needs the opening balance, all incoming transfers, agent cash-in/out, "
      "fees and reversals, which this dataset does not contain. An account can send 1,050,000 RWF it received minutes "
      "earlier.")
    h("4.2 Q2 Daily activity", 2)
    p("30 CAT days; both engines identical; sums reconcile with capture.py (20,003,300 transfers, "
      "231,482,327,870 RWF). Activity runs from 06:00 to 21:59:59 CAT each day.")
    shot("4.2_q2_results", "Q2 daily counts and totals with reconciliation checks.")
    img(CHARTS / "q2_daily.png", "Q2 daily transfer count (thousands) and value (billion RWF).")
    p("Pattern: counts are almost flat (≈666,667 per day), but 6–15 September carry 150–330 extra transfers per day "
      "and noticeably higher value (≈7.74–7.86 bn RWF vs ≈7.66–7.70 bn). Transfer value is not provider revenue: "
      "revenue comes from fees and commissions that are not recorded, many P2P transfers may be free, and money "
      "passing through several wallets is counted at every hop.")
    p("Implementation note: the literal per-row Cypher (datetime() on each of 20 million relationships) took ≈190 s; "
      "grouping by the CAT day number first and converting only the 30 keys returns the same rows in 8.8–40 s.",
      italic=True)
    h("4.3 Q3 Shared devices", 2)
    p("Rows per set (identical in both engines): A 4, B 9, C 9, D 4, E 9, F 9.")
    shot("4.3_q3_0", "Q3 overlapping device use, sets A–C.")
    shot("4.3_q3_1", "Q3 sets D–F and summary.")
    p("Legitimate sharing: in set B, accounts 147226, 147227, 147229 and 147230 share handset 147226 with 147228 for "
      "the whole month, consistent with a household phone; in set A, 220998 shares handset 998 because it is a second "
      "wallet of the same customer. Overlap does not prove simultaneous use: intervals [2,7] and [5,11] overlap on days "
      "5–7, but the accounts could have used the phone at different times on those days, and observations do not say "
      "which device initiated a payment.")
    doc.add_page_break()


def section5():
    h("5 Connected accounts", 1)
    h("5.1 Q4 Ownership and device network", 2)
    q4 = json.loads((RES / "section5" / "q4.json").read_text(encoding="utf-8"))
    rows = []
    for s in "ABCDEF":
        r = {x["depth"]: x["postgres_total"] for x in q4 if x["set"] == s}
        rows.append([s, f"{r[4]:,}", f"{r[6]:,}", f"{r[8]:,}", f"{r[10]:,}"])
    table(["Set", "4 links", "6 links", "8 links", "10 links"], rows)
    p("24/24 comparisons identical (counts by owner region in the screenshots).")
    todo("Write your prediction for depths 8 and 10, made from the 4 and 6 results, and whether it held.")
    shot("5.1_q4", "Q4 queries and the 24 comparisons.")
    shot("5.1_q4_1", "Q4 comparisons (continued) and supporting paths.")
    img(CHARTS / "q4_depth.png", "Q4 distinct accounts reached against depth (log scale).")
    p("Deeper links create weak associations: each hop adds accounts linked only through someone else's phone, an old "
      "replacement handset or a public terminal. C and F reach a public terminal (220020, 220170) within 4 links, so "
      "they grow to ≈1,430 accounts at 10 links; the 6-link path from 998 reaches account 101752 via a household phone "
      "and a replacement handset, with no direct relationship to 998.")
    h("5.2 Q5 Rapid onward payments", 2)
    p("Matches: A 1, B 1, C 1, D 0, E 0, F 0. Example (A): 20000001 998→8917 10:00:00 200,000 RWF; 20000002 "
      "8917→16836 10:00:40 196,000 (98%); 20000003 16836→24755 10:01:20 192,080 (98%).")
    shot("5.2_q5", "Q5 pair, six counts and the set A chain with IDs, CAT times and amounts.")
    p("This does not prove that the same funds moved: wallet money is fungible, 8917 may have paid from an existing "
      "balance, and similar amounts 40 s apart can be coincidence or pre-arranged payments. Balances and the holders' "
      "explanations would be needed.")
    h("5.3 Q6 Circular payments", 2)
    p("Matches: A 1 (4 transfers), C 1 (6 transfers), others 0. Set A: 20000001→20000002→20000003→20000004 back to "
      "998 at 10:02:00, 120 s span, each amount 98% of the previous.")
    shot("5.3_q6", "Q6 pair, six counts and the set A cycle.")
    p("Further evidence: KYC relationships between the accounts, balances before and after, the device that initiated "
      "each payment, similar cycles on other days, and any later cash-out at an agent.")
    h("5.4 Q7 Shortest transfer connection", 2)
    table(["Set", "Review → destination", "≤ 4", "≤ 6"], [
        ["A", "998 → 24755", 3, 3], ["B", "147228 → 186823", "none", 5], ["C", "206978 → 6573", "none", 5],
        ["D", "237998 → 21755", 3, 3], ["E", "144228 → 183823", "none", 5], ["F", "203978 → 3573", "none", 5]])
    shot("5.4_q7", "Q7 pair and the 12 length comparisons.")
    shot("5.4_q7_1", "Q7 supporting paths verified in PostgreSQL (IDs, direction, roles, existence, length).")
    p("No route within four transfers is not evidence of no connection: B, C, E and F connect in five, and accounts "
      "may also be linked via merchants/agents (excluded), cash, other providers or other months. A structural route "
      "need not show onward movement: the B route's transfers fall on 18, 23, 29, 25 and 13 September and range from "
      "4,700 to 968,486 RWF, so money from 147228 cannot have flowed along it.")
    doc.add_page_break()


def section6():
    h("6 Correctness and performance", 1)
    h("6.1 One comparison program", 2)
    p("compare_61.py uses run_pg(query, params) and run_neo(query, params) over persistent connections, normalizes "
      "integers, sorts rows and compares exactly. **61/61 cases match** (Q1 6, Q2 1, Q3 6, Q4 24, Q5 6, Q6 6, Q7 12); "
      "Q7 lengths are compared separately and all 8 supporting paths are valid. Independent Python checks on the CSVs: "
      "the Q3 overlap of 998 and 8917 on device 228001 = [0,29]; the Q6 cycle 20000001–20000004 satisfies every rule. "
      "No mismatches had to be resolved.")
    shot("6.1_cases", "61-case status table.")
    shot("6.1_compare", "Per-question summary and the two independent CSV checks.")

    h("6.2 Timing a fixed selection", 2)
    p("One service at a time, persistent connections, fixed memory settings; 3 untimed warm-ups (180 s) and 5 measured "
      "runs (60 s) per case, perf_counter from submission to the final row; errors and timeouts recorded separately.")
    summ = read_csv(RES / "benchmark" / "summary.csv")
    rows = []
    for r in summ:
        if r["pg_sql"] == "bidirectional":
            continue
        pg = (f"{ms(r['pg_median_ms'])} ({ms(r['pg_min_ms'])}–{ms(r['pg_max_ms'])})" if r["pg_ok"] not in ("", "0")
              else f"{r['pg_timeout']}/5 timeouts")
        neo = f"{ms(r['neo_median_ms'])} ({ms(r['neo_min_ms'])}–{ms(r['neo_max_ms'])})"
        ratio = f"{float(r['pg_neo_median_ratio']):.2f}" if r["pg_neo_median_ratio"] else "–"
        rows.append([r["case"], pg, neo, ratio])
    table(["Case", "PostgreSQL median (min–max) ms", "Neo4j median (min–max) ms", "PG/Neo"], rows,
          widths=[2.6, 5.6, 5.6, 2.2], size=8.5)
    shot("6.2_pg_bench", "PostgreSQL benchmark run (Neo4j stopped).")
    shot("6.2_neo_bench", "Neo4j benchmark run (PostgreSQL stopped).")
    shot("6.2_summary", "Benchmark summary with PG/Neo median ratios; all completed runs correct.")
    shot("6.2_pg_plans", "EXPLAIN (ANALYZE, BUFFERS) for Q2, Q4 B depth 10 and Q7 B bound 6.")
    shot("6.2_neo_plans", "Neo4j PROFILE for Q2, Q4 B depth 10 and Q7 B bound 6.")
    p("Processing difference (Q7 B, bound 6): the PostgreSQL recursive CTE (107.5 s without a timeout) expanded "
      "427,953 (account, depth) states over 7 iterations: 187,957 index-only scans returned 14.6 million transfer rows "
      "and 14.6 million account lookups checked the person role (60 million buffer hits). UNION only removes duplicate "
      "(account, depth) pairs, so accounts are re-expanded at every depth, and the search continues after reaching the "
      "destination. Neo4j used one ShortestPath operator (breadth-first, person filter applied during the search, stops "
      "at the first hit): 18,794 db hits. For Q2 both do full scans: PostgreSQL a Parallel Seq Scan with 2 workers and "
      "HashAggregate (12.1 s), Neo4j a relationship-type scan with 60 million db hits (8.6 s).")

    h("6.3 One SQL improvement", 2)
    p("Prediction: searching from both ends should reduce the work from about d^6 to 2·d^3 states (d ≈ 16 person "
      "contacts per account), i.e. by hundreds of times. Implementation: forward ⌈bound/2⌉ steps from the review "
      "account and backward ⌊bound/2⌋ steps from the destination (incoming transfers from person senders), minimum "
      "depth per account on each side, join on shared accounts, MIN(forward + backward). Both versions are kept.")
    rows = []
    for case in ("Q7-B-b4", "Q7-B-b6", "Q7-E-b4", "Q7-E-b6"):
        plain = next(r for r in summ if r["case"] == case and r["pg_sql"] == "plain")
        bi = next(r for r in summ if r["case"] == case and r["pg_sql"] == "bidirectional")
        rows.append([case, ms(plain["pg_median_ms"]) if plain["pg_ok"] != "0" else "5/5 timeouts",
                     ms(bi["pg_median_ms"]), ms(bi["neo_median_ms"]), f"{float(bi['pg_neo_median_ratio']):.2f}"])
    table(["Case", "Plain SQL ms", "Bidirectional SQL ms", "Neo4j ms", "New PG/Neo"], rows)
    shot("6.3_pg_revised", "Four revised SQL cases timed (B/E at both bounds).")
    shot("6.3_recheck", "All 61 cases recheck correctly with the bidirectional Q7.")
    p("Outcome: the bidirectional version is selected for Section 7. B/E bound-6 medians fell from >60 s (timeouts) "
      "to ≈80 ms, more than 700× faster; plan execution fell from 107.5 s to 0.21 s. Neo4j remains ≈6–8× faster at "
      "bound 6, but the gap is now milliseconds, so most of the original difference came from the SQL algorithm, not "
      "the relational model.")
    doc.add_page_break()


def section7():
    h("7 Concurrent requests", 1)
    h("7.1 Load test", 2)
    p("loadtest.py starts 1, 4, 8 or 16 processes, each with one persistent connection; each warms up A-B-C once, "
      "waits at a Barrier, then runs A-B-C cycles until 15 s after the common start, finishing any started cycle. "
      "Q4 depth 10 and Q7 bound 6 (bidirectional SQL) were tested separately, three repetitions each; correctness "
      "checked after timing. Values below are medians of the three repetitions; there were no errors, timeouts or "
      "wrong answers.")
    ls = read_csv(RES / "loadtest" / "summary.csv")
    rows = [[("PostgreSQL" if r["engine"] == "pg" else "Neo4j"), r["workload"], r["clients"],
             f"{float(r['throughput_rps']):,.1f}", f"{float(r['median_ms']):,.1f}", f"{float(r['p95_ms']):,.1f}",
             f"{float(r['cpu_avg_pct']):.0f}"] for r in ls]
    table(["Engine", "Workload", "Clients", "Req/s", "Median ms", "p95 ms", "CPU avg %"], rows, size=8.5)
    shot("7.1_pg_load", "PostgreSQL load test: 24 repetitions.")
    shot("7.1_neo_load", "Neo4j load test: 24 repetitions.")
    shot("7.1_pg_active_clients", "Active client processes during a 16-client test (Task Manager).")
    shot("7.1_neo_active_clients", "Active client processes during a Neo4j 16-client test.")
    shot("7.1_summary", "All 16 engine/workload/client-count summaries.")
    h("7.2 Scaling", 2)
    img(CHARTS / "loadtest.png", "Throughput and p95 latency against client count.")
    p("Both engines scale almost linearly from 1 to 4 clients, then flatten: average CPU rises from 9–30% at 4 clients "
      "to 96–99% at 8 and 16, so the 8 hardware threads (shared by clients and database) are saturated. From 8 to 16 "
      "clients throughput stays flat or falls slightly while p95 roughly doubles (PostgreSQL Q4 211→464 ms; Neo4j Q7 "
      "17→72 ms): extra clients only queue. Throughput is total completed work per second; response time is what one "
      "analyst waits. At 16 clients PostgreSQL still completes ≈116 Q4 requests/s, but an analyst's p95 wait is 0.46 s "
      "instead of 0.14 s. In this experiment the graph advantage survives the SQL improvement: at 8 clients Neo4j "
      "handled ≈12× (Q4) and ≈30× (Q7) more requests per second. Limits: one laptop, clients and server on the same "
      "CPU, three accounts, 15 s windows, warm caches.")
    doc.add_page_break()


def section8():
    h("8 Decisions", 1)
    h("8.1 Investigation priority", 2)
    table(["Pair", "Q3 devices", "Q4 d10", "Q5 chains", "Q6 cycles", "Q7 ≤4 / ≤6", "Review first"], [
        ["A 998 / D 237998", "4 / 4", "187 / 192", "1 / 0", "1 (4 hops) / 0", "3,3 / 3,3", "A"],
        ["B 147228 / E 144228", "9 / 9", "228 / 230", "1 / 0", "0 / 0", "–,5 / –,5", "B"],
        ["C 206978 / F 203978", "9 / 9", "1,433 / 1,426", "1 / 0", "1 (6 hops) / 0", "–,5 / –,5", "C"],
    ], size=8.5)
    p("Within each pair, device sharing, network size and shortest routes are almost identical; the difference is the "
      "money pattern. A, B and C each show a rapid onward chain within their review hour (A and C also a cycle); D, E "
      "and F show none. D/E/F still reach their destinations, but through payments days apart (D: 45,000 RWF on "
      "6, 8 and 10 September).")
    p("A vs D in detail: on 6 September at 10:00:00, 998 sent 200,000 RWF to 8917; 40 s later 8917 sent 196,000 to "
      "16836; at 10:01:20 16836 sent 192,080 to 24755; at 10:02:00 24755 returned 188,238 to 998. All four person "
      "accounts were observed on handset 228001 for the whole month. D (237998) has a similar device network, but its "
      "45,000 RWF payments to 5917 → 13836 → 21755 are two days apart and never return.")
    img(FIGS / "cycle_A.png", "Transfer cycle for set A (Q6).")
    img(FIGS / "network_A.png", "Ownership and device links around account 998 (Q3/Q4).")
    p("Ordinary explanation: four relatives or friends sharing one phone could be passing a savings-group (ikimina) "
      "contribution around, each keeping 2%, or testing a new wallet. Missing evidence: KYC relationships, balances, the "
      "device that initiated each payment, other months, and any cash-out at an agent. A documented savings group would "
      "lower the priority; a quick cash-out after the cycle would raise it. These records justify enquiries, not a "
      "finding of fraud.")

    h("8.2 Model recommendation", 2)
    todo("Edit this ≤500-word recommendation into your own words; keep the evidence.")
    for para in [
        "**Reporting → PostgreSQL.** Q1 hourly activity took 0.5–1.2 ms in PostgreSQL vs ≈8 ms in Neo4j. Daily totals "
        "(Q2) are a natural SQL GROUP BY with time-zone functions. Evidence against: Neo4j scanned the month faster "
        "(8.8 s vs 12.5 s), but only after rewriting the Cypher (the literal per-row form took ≈190 s).",
        "**Investigations → Neo4j.** Device networks (Q4), cycles (Q6) and shortest routes (Q7) are path problems. "
        "Cypher states them as patterns with predicates inside the search; SQL needs long recursive CTEs. Neo4j stayed "
        "at 4–14 ms for Q4 depth 10 and Q7 bound 6, where PostgreSQL took 23 ms and 80 ms (after the rewrite) or timed "
        "out (before).",
        "**Concurrency.** At 8–16 clients on this laptop Neo4j delivered 12–32× the throughput with 13–23× lower p95; "
        "both were CPU-bound from 8 clients.",
        "**Load time and storage.** PostgreSQL: 34.6 s load + 96.8 s keys/indexes, 3.73 GB. Neo4j: 58.1 s offline "
        "import + 31.7 s constraints/index, 2.05 GB store. The covering transfer indexes (≈2 GB) explain most of "
        "PostgreSQL's size.",
        "**Plan comparison.** For Q7 the recursive CTE re-expanded 14.6 million transfer rows; Neo4j's ShortestPath "
        "operator needed 18,794 db hits. The graph model's benefit is stored adjacency and readable path patterns; the "
        "rest was algorithm: Neo4j's breadth-first search with early stop versus a CTE without a visited set. Our "
        "bidirectional SQL recovered more than 700× without changing the schema.",
        "**Freshness.** Keep PostgreSQL as the system of record and copy to Neo4j for investigations, at most 15 minutes "
        "behind (change data capture or micro-batches). Stale data would miss a cycle completed in the last minutes or "
        "show closed accounts: acceptable for retrospective investigation, not for real-time blocking.",
        "**Limits.** (1) One laptop, clients and server sharing 8 threads, warm caches, one month of synthetic data. "
        "(2) Six review accounts and no concurrent writes: production load and data growth could change the ranking.",
    ]:
        p(para)

    h("8.3 Individual judgement", 2)
    for name, sid in (("BIZIMUNGU Theoneste", "101398"), ("NIYITANGA Eric", "101393"), ("NDASHIMYE Janvier", "101363")):
        p(f"{name} ({sid})", bold=True)
        todo("150–200 words in your own words: one measured finding, its explanation, and a check that could "
             "overturn it.")
    h("Contributions", 3)
    table(["Member", "Contribution"], [["BIZIMUNGU Theoneste", ""], ["NIYITANGA Eric", ""], ["NDASHIMYE Janvier", ""]],
          widths=[5, 12])
    todo("Fill in who did what.")
    h("LLM and external assistance", 3)
    p("Claude (Anthropic, Claude Opus 5.5, via Claude Code) helped write and debug the Python, SQL and Cypher code, the "
      "benchmark and load-test programs and the diagrams, and drafted explanation text. We checked it by: comparing "
      "every query pair across both engines (61 cases); recomputing one device overlap and one cycle directly from the "
      "CSV files in plain Python; verifying every Q7 path in PostgreSQL; reading the execution plans; and running all "
      "scripts ourselves on our own machine. All explanations were reviewed and rewritten by the group.")
    todo("Adjust this statement so it accurately describes how your group used assistance and checked it.")
    doc.add_page_break()


def section9():
    h("9 Reproducible work", 1)
    p("GitHub: https://github.com/bibabbly/auca-abda-p2-mobile-money (the ZIP matches the final commit). The README "
      "lists all commands in execution order; .env.example holds placeholder secrets; captured CSVs and database stores "
      "stay local. The generator is unchanged (SHA-256 2ea0e59d2b7b987d84a3d151bd9d1ad7b3052bfbc25abf1466c8e94467e9455a).")
    shot("9.2_final_compare", "Final rerun of the 61-case comparison.")
    shot("9.2_services_stopped", "PostgreSQL and Neo4j services stopped; stores retained.")


# ------------------------------------------------------------------ build
setup()
cover()
section1(); section2(); section3(); section4(); section5(); section6(); section7(); section8(); section9()
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Saved {OUT}")
print(f"Figures: {fig_no[0]}; missing screenshots marked in red: {missing or 'none'}")
